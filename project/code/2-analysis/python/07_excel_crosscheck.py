"""Langkah 7b: cek silang dengan Excel. Data bersih ditulis ke crosscheck.xlsx, defect rate dan z-test dihitung dengan RUMUS Excel
(SUMIFS, COUNTIFS, NORM.S.DIST), dihitung ulang oleh LibreOffice, lalu dibandingkan dengan hasil Python. Semua baris harus PASS."""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import common as C  # noqa: E402
from kpi_pandas import prepare  # noqa: E402

prod, ng, op, cm = prepare()
Z = pd.read_csv(C.RES_ANALYSIS / "ztest.csv").set_index("scope")
wb = Workbook()
ws = wb.active; ws.title = "production"
ws.append(["date", "shift", "harness_type", "operator_id", "input_qty", "ok_qty", "ng_qty", "period"])
for i, r in enumerate(prod.itertuples(), start=2):
    ws.append([r.date.to_pydatetime(), r.shift, r.harness_type, r.operator_id, r.input_qty, r.ok_qty, r.ng_qty,
               f'=IF(A{i}<DATE(2026,7,1),"before",IF(A{i}>=DATE(2026,8,1),"after","transisi"))'])
wn = wb.create_sheet("ng_log")
wn.append(["ng_id", "date", "shift", "harness_type", "operator_id", "station_found", "defect_code", "cause_source", "bracket_source", "period", "is_id"])
for i, r in enumerate(ng.itertuples(), start=2):
    wn.append([r.ng_id, r.date.to_pydatetime(), r.shift, r.harness_type, r.operator_id, r.station_found, r.defect_code, r.cause_source, r.bracket_source,
               f'=IF(B{i}<DATE(2026,7,1),"before",IF(B{i}>=DATE(2026,8,1),"after","transisi"))', f'=IF(LEFT(G{i},3)="ID-",1,0)'])
for s in (ws, wn):
    for c in s[1]:
        c.font = Font(bold=True)
    for r in s.iter_rows(min_row=2, min_col=1 if s is ws else 2, max_col=1 if s is ws else 2):
        r[0].number_format = "yyyy-mm-dd"
nP, nN = len(prod) + 1, len(ng) + 1
wc = wb.create_sheet("check", 0)
hdr = ["scope", "n_before", "x_before", "p_before", "n_after", "x_after", "p_after", "diff", "p_pool", "se_pool", "z", "p_value", "se_wald", "ci_low", "ci_high"]
wc.append(["Cek silang: semua sel di bawah adalah RUMUS Excel dari sheet production dan ng_log"]); wc["A1"].font = Font(bold=True)
wc.append(hdr)
for c in wc[2]:
    c.font = Font(bold=True)
SC = {  # scope -> (filter di production, filter di ng_log)
    "Insulation damage (ID-01+02+03)": ("", ""),
    **{f"Harness tipe {t}": (f',production!$C$2:$C${nP},"{t}"', f',ng_log!$D$2:$D${nN},"{t}"') for t in "ABC"},
    **{s: (f',production!$B$2:$B${nP},"{s}"', f',ng_log!$C$2:$C${nN},"{s}"') for s in C.SHIFTS},
}
for k, (fp, fn) in SC.items():
    r = wc.max_row + 1
    pj = lambda per: f'SUMIFS(production!$E$2:$E${nP},production!$H$2:$H${nP},"{per}"{fp})'
    nj = lambda per: f'SUMIFS(ng_log!$K$2:$K${nN},ng_log!$J$2:$J${nN},"{per}"{fn})'
    wc.append([k, f"={pj('before')}", f"={nj('before')}", f"=C{r}/B{r}", f"={pj('after')}", f"={nj('after')}", f"=F{r}/E{r}", f"=D{r}-G{r}",
               f"=(C{r}+F{r})/(B{r}+E{r})", f"=SQRT(I{r}*(1-I{r})*(1/B{r}+1/E{r}))", f"=H{r}/J{r}", f"=2*(1-_xlfn.NORM.S.DIST(ABS(K{r}),TRUE))",
               f"=SQRT(D{r}*(1-D{r})/B{r}+G{r}*(1-G{r})/E{r})", f"=H{r}-1.959963984540054*M{r}", f"=H{r}+1.959963984540054*M{r}"])
r0 = wc.max_row + 2
wc.cell(r0, 1, "Per sumber penyebab (COUNTIFS)").font = Font(bold=True)
causes = ["ID-01", "ID-02", "ID-03"]
for j, code in enumerate(causes):
    r = r0 + 1 + j
    wc.append([code, f"=SUMIFS(production!$E$2:$E${nP},production!$H$2:$H${nP},\"before\")", f'=COUNTIFS(ng_log!$G$2:$G${nN},"{code}",ng_log!$J$2:$J${nN},"before")', f"=C{r}/B{r}",
               f"=SUMIFS(production!$E$2:$E${nP},production!$H$2:$H${nP},\"after\")", f'=COUNTIFS(ng_log!$G$2:$G${nN},"{code}",ng_log!$J$2:$J${nN},"after")', f"=F{r}/E{r}", f"=D{r}-G{r}"])
wc.column_dimensions["A"].width = 36
tmp = Path(tempfile.mkdtemp()); src = tmp / "in" / "crosscheck.xlsx"; src.parent.mkdir()
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)], check=True, capture_output=True, timeout=240)
shutil.copy(tmp / "out" / "crosscheck.xlsx", C.RES_ANALYSIS / "crosscheck.xlsx")
v = load_workbook(C.RES_ANALYSIS / "crosscheck.xlsx", data_only=True)["check"]
rows = []
cmp = lambda scope, metric, xl, py, tol: rows.append((scope, metric, xl, py, "PASS" if xl is not None and abs(xl - py) <= tol else "FAIL"))
for r in v.iter_rows(min_row=3, max_row=2 + len(SC), values_only=True):
    z = Z.loc[r[0]]
    cmp(r[0], "input before", r[1], z.n_before, 0); cmp(r[0], "NG before", r[2], z.x_before, 0); cmp(r[0], "input after", r[4], z.n_after, 0); cmp(r[0], "NG after", r[5], z.x_after, 0)
    cmp(r[0], "rate before", r[3], z.rate_before, 1e-12); cmp(r[0], "rate after", r[6], z.rate_after, 1e-12); cmp(r[0], "z", r[10], z.z, 1e-6)
    cmp(r[0], "p-value", r[11], z.p_value, 1e-9); cmp(r[0], "CI low", r[13], z.ci_low, 1e-9); cmp(r[0], "CI high", r[14], z.ci_high, 1e-9)
cs = pd.read_csv(C.RES_ANALYSIS / "cause.csv")
for i, code in enumerate(causes):
    r = next(v.iter_rows(min_row=r0 + 1 + i, max_row=r0 + 1 + i, values_only=True))
    cmp(code, "rate before", r[3], cs[(cs.defect_code == code) & (cs.period == "before")].rate.iloc[0], 1e-12)
    cmp(code, "rate after", r[6], cs[(cs.defect_code == code) & (cs.period == "after")].rate.iloc[0], 1e-12)
df = pd.DataFrame(rows, columns=["scope", "metric", "excel", "python", "status"]); df.to_csv(C.RES_ANALYSIS / "crosscheck_excel.csv", index=False)
print(f"{(df.status == 'PASS').sum()} dari {len(df)} PASS")
assert (df.status == "PASS").all(), df[df.status != "PASS"]
